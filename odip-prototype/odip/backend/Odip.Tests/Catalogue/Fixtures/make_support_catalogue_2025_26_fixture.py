"""Builds support-catalogue-2025-26-trimmed.xlsx, the 2025-26 test fixture for the catalogue importer.

The official file is the NDIA "Support Catalogue 2025-26" workbook (effective 1 July 2025, about 550 KB because it carries
two hidden working sheets). It is a public document and is NOT committed; this script cuts it down to what the importer
tests need while keeping the real layout, so the 2025-26 (per-state price columns) path is tested against real rows:

  * only the two visible sheets are kept ("Current Support Items" and "Legacy Support Items"; the hidden "Sheet3" and
    "Support Catalogue" working sheets are dropped, which is also what the importer does);
  * every row of support categories 01 and 04 is kept (the Oassist families and everything that shares a code prefix
    with them), plus a sample of each other category and every row that exercises a parsing edge: the trailing
    "Bereavement" row (a code with no underscores and blank claim flags), rows whose start date is stored as text, rows
    whose claim flags are "N/A" or 0, Quotable and "Unit Price = $1" rows;
  * the whole Legacy sheet (31 rows) is kept;
  * cell values and number formats are copied as they are, so numeric dates, text dates and "$" formats survive.

Usage (needs openpyxl, no network):

    python make_support_catalogue_2025_26_fixture.py --source <official 2025-26 .xlsx> [--output <fixture .xlsx>]

The official file is https://www.ndis.gov.au/media/7727/download?attachment (fetched 2026-10-02). The output must stay
under 100 KB; the script fails if it does not, and Odip.Tests/Catalogue/CatalogueFixtureTests.cs checks it again.
"""
import argparse
import os
import sys

import openpyxl

CURRENT = "Current Support Items"
LEGACY = "Legacy Support Items"
KEEP_CATEGORIES = ("01", "04")
SAMPLE_PER_OTHER_CATEGORY = 3
MAX_BYTES = 100 * 1024


def header_index(ws):
    return {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value is not None}


def is_edge_row(row, cols):
    code = str(row[cols["Support Item Number"]].value or "")
    if code.count("_") < 4:  # "Bereavement": no structured code
        return True
    start = row[cols["Start date"]].value
    if isinstance(start, str):  # a few rows store the start date as text
        return True
    for name in ("Non-Face-to-Face Support Provision", "Provider Travel", "Short Notice Cancellations."):
        if str(row[cols[name]].value).strip() in ("N/A", "0"):
            return True
    return False


def select_current_rows(ws):
    cols = header_index(ws)
    kept, per_category, seen_types = [], {}, set()
    for row in ws.iter_rows(min_row=2):
        code = str(row[cols["Support Item Number"]].value or "")
        category = code[:2]
        type_value = row[cols["Type"]].value
        keep = category in KEEP_CATEGORIES or is_edge_row(row, cols)
        if not keep:
            key = (category, type_value)
            if per_category.get(category, 0) < SAMPLE_PER_OTHER_CATEGORY or key not in seen_types:
                keep = True
                seen_types.add(key)
        if keep:
            per_category[category] = per_category.get(category, 0) + 1
            kept.append(row)
    return kept


def copy_sheet(src, dst, rows):
    for c in src[1]:
        out = dst.cell(row=1, column=c.column, value=c.value)
        out.number_format = c.number_format
    for r, row in enumerate(rows, start=2):
        for c in row:
            out = dst.cell(row=r, column=c.column, value=c.value)
            out.number_format = c.number_format


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, help="official 2025-26 support catalogue .xlsx")
    parser.add_argument("--output", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "support-catalogue-2025-26-trimmed.xlsx"))
    args = parser.parse_args()

    src = openpyxl.load_workbook(args.source)
    out = openpyxl.Workbook()
    out.remove(out.active)

    current_rows = select_current_rows(src[CURRENT])
    legacy_rows = list(src[LEGACY].iter_rows(min_row=2))
    copy_sheet(src[CURRENT], out.create_sheet(CURRENT), current_rows)
    copy_sheet(src[LEGACY], out.create_sheet(LEGACY), legacy_rows)
    out.save(args.output)

    size = os.path.getsize(args.output)
    print(f"{args.output}: {len(current_rows)} current rows, {len(legacy_rows)} legacy rows, {size} bytes")
    if size >= MAX_BYTES:
        sys.exit(f"fixture is {size} bytes; the limit is {MAX_BYTES}")


if __name__ == "__main__":
    main()
